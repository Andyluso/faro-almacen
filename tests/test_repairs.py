import io
import os
import sqlite3
import tempfile
import zipfile
from pathlib import Path
from unittest.mock import patch

os.environ['SUPABASE_URL'] = ''
os.environ['SUPABASE_KEY'] = ''
os.environ['SUPABASE_SERVICE_ROLE_KEY'] = ''
os.environ['FARO_PRIVATE_CLOUD_READY'] = '0'
os.environ['FARO_USERNAME'] = 'tester'
os.environ['FARO_PASSWORD'] = 'test-password-12345'

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from backend import database as db
from backend import reliability

ROOT = Path(__file__).resolve().parents[1]

class SessionClient(TestClient):
    """Keep legacy regression cases while exercising the new cookie login."""
    def request(self, method, url, **kwargs):
        auth = kwargs.pop('auth', None)
        self.cookies.clear()
        if isinstance(auth, tuple):
            login = super().request('POST','/api/auth/login',json={'username':auth[0],'password':auth[1]})
            if login.status_code != 200: return login
            headers = dict(kwargs.pop('headers', {}) or {})
            headers['X-CSRF-Token'] = login.json()['csrf']
            kwargs['headers'] = headers
        return super().request(method,url,**kwargs)

@pytest.fixture()
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(db, 'DB_PATH', str(tmp_path / 'inventario.db'))
    # Import main only after replacing the database path; tests cannot touch real data.
    from backend import main
    db.init_db()
    from backend.team_store import initialize
    initialize()
    monkeypatch.setattr(main, 'UPLOADS_DIR', str(tmp_path / 'uploads'))
    monkeypatch.setattr(main, 'EXCELS_DIR', str(tmp_path / 'uploads' / 'excels'))
    Path(main.EXCELS_DIR).mkdir(parents=True)
    main.app.middleware_stack = None
    return SessionClient(main.app)

AUTH = ('tester', 'test-password-12345')

@pytest.mark.parametrize('path', ['/', '/api/audits', '/api/employees', '/api/tasks', '/api/system/backup', '/uploads/test.xlsx', '/docs'])
def test_every_private_surface_requires_login(client, path):
    response=client.get(path)
    assert response.status_code == (401 if path.startswith('/api/') else 200)
    if not path.startswith('/api/'): assert 'Ingresa a FARO' in response.text
    assert 'www-authenticate' not in response.headers

def test_auth_and_origin(client):
    assert client.get('/api/audits', auth=AUTH).status_code == 200
    assert client.get('/api/audits', auth=('tester', 'wrong')).status_code == 401
    assert client.post('/api/tasks', auth=AUTH, headers={'Origin':'https://attacker.invalid'}, json={'title':'bad'}).status_code == 403
    assert client.post('/api/team/agenda', auth=AUTH, json={'title':'Tarea válida','start_date':'2026-09-25'}).status_code == 200
    assert client.get('/api/tasks', auth=AUTH).headers['cache-control'] == 'no-store'

def test_rate_limit(client):
    for _ in range(10):
        client.get('/api/audits', auth=('tester', 'wrong'))
    assert client.get('/api/audits', auth=('tester', 'wrong')).status_code == 429

def test_health_and_cache_cleanup(client):
    assert client.get('/healthz').status_code == 200
    response = client.get('/sw.js')
    assert response.status_code == 200
    assert 'caches.delete' in response.text
    assert 'caches.put' not in response.text

def test_only_one_truthful_cloud_status_route(client):
    from backend.main import app
    assert sum(getattr(r, 'path', '') == '/api/system/database-status' for r in app.routes) == 1
    data = client.get('/api/system/database-status', auth=AUTH).json()
    assert data['supabase_connected'] is False
    assert data['mode'] == 'sqlite'

def test_connection_is_actually_probed(client):
    with patch.dict(os.environ, {'SUPABASE_URL':'https://example.invalid'}):
        with patch('backend.cloud_adapter.get_client') as get_client:
            get_client.return_value.table.return_value.select.return_value.limit.return_value.execute.side_effect = TimeoutError()
            assert reliability.status(probe=True)['cloud_reachable'] is False
            get_client.return_value.table.return_value.select.return_value.limit.return_value.execute.assert_called_once()

def test_employee_delete_removes_shifts(client):
    emp = db.create_employee('Prueba')
    schedule = db.get_or_create_weekly_schedule('2026-09-21')
    db.save_schedule_shifts(schedule['id'], [{'employee_id':emp['id'],'day_of_week':'Lun','shift_type':'Libre'}])
    assert db.delete_employee(emp['id'])
    with db.get_connection() as conn:
        assert conn.execute('SELECT COUNT(*) FROM schedule_shifts WHERE employee_id=?', (emp['id'],)).fetchone()[0] == 0

def test_transactional_change_tracking(client):
    with db.get_connection() as conn:
        before = conn.execute('SELECT COUNT(*) FROM sync_outbox').fetchone()[0]
        conn.execute("INSERT INTO store_tasks(title) VALUES ('rolled back')")
        conn.rollback()
        assert conn.execute('SELECT COUNT(*) FROM sync_outbox').fetchone()[0] == before
    db.create_store_task('Durable task')
    with db.get_connection() as conn:
        assert conn.execute("SELECT COUNT(*) FROM sync_outbox WHERE table_name='store_tasks'").fetchone()[0] > 0

def test_backup_contains_all_modules_and_assets(client, tmp_path):
    db.create_store_task('Saved in backup')
    assets = tmp_path / 'uploads'
    (assets/'photo.jpg').write_bytes(b'example')
    response = client.get('/api/system/backup', auth=AUTH)
    assert response.status_code == 200
    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        assert 'backend/uploads/photo.jpg' in archive.namelist()
        assert not any('.env' in name or 'access' in name for name in archive.namelist())
        restored = tmp_path/'restored.db'
        restored.write_bytes(archive.read('backend/inventario.db'))
    with sqlite3.connect(restored) as conn:
        assert conn.execute('PRAGMA quick_check').fetchone()[0] == 'ok'
        assert conn.execute("SELECT COUNT(*) FROM store_tasks WHERE title='Saved in backup'").fetchone()[0] == 1
        tables = {r[0] for r in conn.execute("SELECT name FROM sqlite_master WHERE type='table'")}
        assert set(reliability.TABLES).issubset(tables)

def test_restart_keeps_local_changes(client):
    task = db.create_store_task('Keep after restart')
    with patch('backend.cloud_adapter.get_client', side_effect=AssertionError('No cloud hydration')):
        db.init_db()
    assert any(t['id'] == task['id'] for t in db.list_store_tasks())

def test_no_invented_shifts(client):
    assert db.get_hub_summary()['today_shifts'] == []

def test_excel_import_export_and_validation(client):
    sample = ROOT/'sample_data/inventario_lunes_muestra.xlsx'
    with sample.open('rb') as f:
        result = client.post('/api/audits/upload', auth=AUTH, files={'file':(sample.name,f)}, data={'audit_name':'Prueba'})
    assert result.status_code == 200, result.text
    aid = result.json()['audit_id']
    items = client.get(f'/api/audits/{aid}', auth=AUTH).json()['items']
    assert len(items) >= 200
    iid = items[0]['id']
    result = client.post(f'/api/items/{iid}/validate', auth=AUTH, json={'verdict':'Error de lectura','notes':'Prueba'})
    assert result.status_code == 200
    response = client.get(f'/api/export/{aid}/excel', auth=AUTH)
    assert response.status_code == 200
    wb = load_workbook(io.BytesIO(response.content))
    assert len(wb.sheetnames) > 0

def test_empty_excel_returns_client_error(client):
    wb = Workbook(); stream = io.BytesIO(); wb.save(stream)
    response = client.post('/api/audits/upload', auth=AUTH, files={'file':('empty.xlsx',stream.getvalue())})
    assert response.status_code == 400

def test_invalid_photo_rejected(client):
    response = client.post('/api/references/12345678/photo', auth=AUTH, files={'file':('bad.jpg',b'<script>alert(1)</script>')})
    assert response.status_code == 400

def test_email_export_escapes_stored_content():
    from backend.excel_exporter import generate_email_table_html
    payload = '<img src=x onerror=alert(1)>'
    html = generate_email_table_html({'name':payload}, [{'name':payload,'reference':'123','size':'M','difference':-1,'validation_notes':payload}])
    assert '<img src=x' not in html
    assert '&lt;img' in html

def test_cloud_backup_disabled_by_default(client):
    from backend import private_backup
    with patch('backend.cloud_adapter.get_client', side_effect=AssertionError('Must not contact cloud')):
        assert private_backup.upload_backup() is False

def test_cloud_backup_refuses_public_key(client, monkeypatch):
    from backend import private_backup
    monkeypatch.setenv('FARO_PRIVATE_CLOUD_READY', '1')
    monkeypatch.setenv('SUPABASE_KEY', 'sb_publishable_example')
    with patch('backend.cloud_adapter.get_client', side_effect=AssertionError('Must not contact cloud')):
        assert private_backup.upload_backup() is False

def test_cloud_backup_refuses_public_bucket(client, monkeypatch):
    from backend import private_backup
    monkeypatch.setenv('FARO_PRIVATE_CLOUD_READY', '1')
    monkeypatch.setenv('SUPABASE_SERVICE_ROLE_KEY', 'sb_secret_TEST_ONLY')
    with patch('backend.cloud_adapter.get_client') as get_client:
        get_client.return_value.storage.get_bucket.return_value = {'public': True}
        assert private_backup.upload_backup() is False
        get_client.return_value.storage.from_.assert_not_called()

def test_failed_cloud_backup_keeps_changes_pending(client, monkeypatch):
    from backend import private_backup
    monkeypatch.setenv('FARO_PRIVATE_CLOUD_READY', '1')
    monkeypatch.setenv('SUPABASE_SERVICE_ROLE_KEY', 'sb_secret_TEST_ONLY')
    db.create_store_task('Must survive failure')
    with reliability.connection() as conn:
        before = conn.execute('SELECT COUNT(*) FROM sync_outbox').fetchone()[0]
    with patch('backend.cloud_adapter.get_client') as get_client:
        get_client.return_value.storage.get_bucket.return_value = {'public': False}
        get_client.return_value.storage.from_.return_value.upload.side_effect = TimeoutError()
        assert private_backup.upload_backup() is False
    with reliability.connection() as conn:
        assert conn.execute('SELECT COUNT(*) FROM sync_outbox').fetchone()[0] == before

def test_backup_retry_and_concurrent_changes(client, monkeypatch):
    from backend import private_backup
    monkeypatch.setenv('FARO_PRIVATE_CLOUD_READY', '1')
    monkeypatch.setenv('SUPABASE_SERVICE_ROLE_KEY', 'sb_secret_TEST_ONLY')
    db.create_store_task('Included in snapshot')
    uploaded = []
    def receive(name, content, **kwargs):
        uploaded.append(content)
        db.create_store_task('Changed during upload')
    with patch('backend.cloud_adapter.get_client') as get_client:
        get_client.return_value.storage.get_bucket.return_value = {'public': False}
        get_client.return_value.storage.from_.return_value.upload.side_effect = receive
        assert private_backup.upload_backup() is True
    with reliability.connection() as conn:
        remaining = conn.execute('SELECT payload FROM sync_outbox').fetchall()
        assert len(remaining) == 1
        assert 'Changed during upload' in remaining[0][0]
    with zipfile.ZipFile(io.BytesIO(uploaded[0])) as archive:
        assert 'backend/inventario.db' in archive.namelist()

def test_photo_proxy_rejects_untrusted_destination(client):
    assert client.get('/api/system/photo?url=https://attacker.invalid/image', auth=AUTH).status_code == 404

