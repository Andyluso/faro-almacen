import os
import re
import unicodedata
from typing import List, Dict, Any, Tuple
import openpyxl
import pandas as pd

def normalize_text(text: Any) -> str:
    """Normaliza texto eliminando acentos, espacios extras y convirtiendo a minúsculas."""
    if text is None:
        return ""
    s = str(text).strip().lower()
    # Eliminar acentos
    s = ''.join(c for c in unicodedata.normalize('NFD', s) if unicodedata.category(c) != 'Mn')
    # Reemplazar caracteres no alfanuméricos por guiones bajos
    s = re.sub(r'[^a-z0-9]', '_', s)
    s = re.sub(r'_+', '_', s).strip('_')
    return s

def classify_category(name: str) -> str:
    """Clasifica el tipo de prenda basándose en palabras clave en su nombre o descripción."""
    n = normalize_text(name)
    
    if any(k in n for k in ['camiseta', 'camisilla', 't_shirt', 'polo', 'tank', 'esqueleto']):
        return 'Camisetas'
    if any(k in n for k in ['pantalon', 'jean', 'denim', 'jogger', 'legging', 'cargo', 'drill']):
        return 'Pantalones'
    if any(k in n for k in ['falda', 'skirt', 'minifalda']):
        return 'Faldas'
    if any(k in n for k in ['chaqueta', 'jacket', 'blazer', 'chaleco', 'cazadora', 'bomber', 'parka', 'gabardina']):
        return 'Chaquetas'
    if any(k in n for k in ['short', 'bermuda']):
        return 'Shorts'
    if any(k in n for k in ['buzo', 'hoodie', 'sueter', 'saco', 'cardigan', 'sweater']):
        return 'Buzos'
    if any(k in n for k in ['vestido', 'dress', 'enterizo', 'jumpsuit', 'romper']):
        return 'Vestidos'
    if any(k in n for k in ['camisa', 'blusa', 'crop_top', 'top', 'blouse']):
        return 'Camisas y Blusas'
    if any(k in n for k in ['correa', 'cinturon', 'gorra', 'sombrero', 'bolso', 'morral', 'cartera', 'medias', 'calcetin', 'billetera', 'accesorio']):
        return 'Accesorios'
        
    return 'Otras Prendas'

# ============================================================================
# TABLA MAESTRA DE CODIFICACIÓN OFICIAL SEVEN SEVEN
# Estructura de 8 dígitos: [2 dígitos Género] + [2 dígitos Tipo Prenda] + [4 dígitos Modelo]
# ============================================================================
SEVEN_SEVEN_GARMENT_TYPES = {
    '00': 'Pantaloncillo',
    '01': 'Camisa',
    '06': 'Buzo',
    '07': 'Pantalón',
    '08': 'Chaqueta',
    '09': 'Camiseta',
    '10': 'Bermuda',
    '11': 'Polo',
    '12': 'Blusa',
    '14': 'Falda',
    '15': 'Conjunto 2 Piezas',
    '16': 'Jean',
    '17': 'Vestidos',
    '18': 'Ropa Interior',
    '19': 'Short',
    '20': 'Body',
    '21': 'Accesorios',
    '22': 'Top',
    '23': 'Leggins',
    '25': 'Chaleco',
    '26': 'Capri',
    '27': 'Ropa Playa',
    '29': 'Overol',
    '32': 'Enterizo',
    '33': 'Saco',
    '38': 'Varios',
    '40': 'Blazer',
    '41': 'Pantaloncillo x2',
    '42': 'Pantaloncillo x3',
    '45': 'Anillo',
    '46': 'Aretes',
    '47': 'Bufanda',
    '48': 'Collar',
    '49': 'Gafas',
    '50': 'Pulsera/Tobillera',
    '51': 'Set de Accesorios',
    '52': 'Sombrero/Caps',
    '53': 'Accesorios de Cabello',
    '54': 'Misceláneo',
    '55': 'Colaboraciones',
    '56': 'Autoliquidable',
    '57': 'Accesorios Belleza',
    '58': 'Maquillaje',
    '59': 'Esmalte',
    '60': 'Perfumería',
    '61': 'Corporales',
    '62': 'Bolso/Canguros',
    '63': 'Reloj',
    '64': 'Billetera',
    '98': 'Medias'
}

def decode_seven_seven_reference(ref_val: Any, name_val: Any = "") -> Tuple[str, str, str]:
    """
    Decodifica una referencia oficial de Seven Seven:
    - 28: Dama
    - 45: Caballero (Únicamente el 45 representa Caballero)
    - Dígitos 3 y 4: Código oficial del tipo de prenda según la tabla maestra de la marca (solo aplica para 28 o 45).
    Retorna: (gender, garment_code, garment_type)
    """
    s_ref = str(ref_val or '').strip()
    digits = re.sub(r'\D', '', s_ref)

    gender = "Otro"
    garment_code = ""
    garment_type = ""

    if len(digits) >= 4:
        dept = digits[:2]
        if dept == "28":
            gender = "Dama"
            code = digits[2:4]
            if code in SEVEN_SEVEN_GARMENT_TYPES:
                garment_code = code
                garment_type = SEVEN_SEVEN_GARMENT_TYPES[code]
        elif dept == "45":
            gender = "Caballero"
            code = digits[2:4]
            if code in SEVEN_SEVEN_GARMENT_TYPES:
                garment_code = code
                garment_type = SEVEN_SEVEN_GARMENT_TYPES[code]

    # Respaldo si no es código numérico estándar de Seven Seven
    s_name = str(name_val or '').strip()
    if gender == "Otro" and s_name:
        norm_name = normalize_text(s_name)
        if any(k in norm_name for k in ['dama', 'mujer', 'femenino', 'femenina']):
            gender = "Dama"
        elif any(k in norm_name for k in ['caballero', 'hombre', 'masculino']):
            gender = "Caballero"

    if not garment_type and s_name:
        garment_type = classify_category(s_name)

    return (gender, garment_code, garment_type)


def clean_reference(raw_val: Any) -> str:
    """
    Limpia y normaliza la referencia de material.
    Si la referencia tiene 11 dígitos (como en el ERP/SAP de la tienda),
    elimina los últimos 3 dígitos y conserva los 8 primeros números.
    Ejemplo: '28080984009' -> '28080984'
    """
    if raw_val is None:
        return ""
    
    # Manejo de formatos numéricos que openpyxl o pandas leen como float o int
    if isinstance(raw_val, float):
        if raw_val.is_integer():
            s = str(int(raw_val))
        else:
            s = str(raw_val)
    elif isinstance(raw_val, int):
        s = str(raw_val)
    else:
        s = str(raw_val).strip()
        if s.endswith('.0'):
            s = s[:-2]

    s = s.strip()
    if not s:
        return ""

    # Extraer solo dígitos
    digits_only = re.sub(r'\D', '', s)
    # Si tiene 11 dígitos o más de 8 dígitos numéricos, los 8 primeros son los válidos (se quitan los últimos 3)
    if len(digits_only) >= 11:
        return digits_only[:8]
    elif len(s) > 8 and s.isdigit():
        if len(s) - 3 >= 8:
            return s[:-3]
        return s[:8]

    return s

def clean_barcode(raw_val: Any) -> str:
    """Limpia el código de barras eliminando decimales de flotantes."""
    if raw_val is None:
        return ""
    if isinstance(raw_val, float):
        if raw_val.is_integer():
            return str(int(raw_val))
        return str(raw_val)
    s = str(raw_val).strip()
    if s.endswith('.0'):
        s = s[:-2]
    return s

COLUMN_PATTERNS = {
    'reference': ['material', 'referencia', 'ref', 'materiales', 'cod_material', 'codigo_referencia', 'codigo_ref', 'item', 'articulo', 'cod_articulo', 'mat'],
    'name': ['descripcion', 'nombre', 'desc', 'prenda', 'nombre_prenda', 'descripcion_prenda', 'detalle', 'articulo'],
    'size': ['talla', 'size', 'tamano', 'medida'],
    'color': ['color', 'col', 'descripcion_color', 'tono'],
    'barcode': ['ean', 'codigo_barras', 'codigo_de_barras', 'barcode', 'ean13', 'cod_barras', 'plu', 'upc', 'codigo'],
    'store_count': ['piso_de_venta', 'tienda', 'piso_venta', 'piso', 'lectura_tienda', 'conteo_tienda', 'fisico_tienda', 'leido_tienda', 'cant_tienda', 'tienda_fisico', 'venta'],
    'warehouse_count': ['bodega', 'lectura_bodega', 'conteo_bodega', 'fisico_bodega', 'leido_bodega', 'cant_bodega', 'bodega_fisico'],
    'theoretical_count': ['teorico', 'teoria', 'stock_teorico', 'sistema', 'stock_sistema', 'esperado', 'saldo_sistema', 'teorica'],
    'difference': ['diferencia', 'diferencias', 'dif', 'sobrante_faltante', 'variacion', 'ajuste', 'saldo_diferencia'],
    'category': ['categoria', 'cat', 'subcategoria', 'linea', 'familia', 'tipo'],
    'observation': ['observacion', 'observaciones', 'obs', 'motivo']
}

def map_headers(raw_headers: List[str]) -> Dict[str, int]:
    """
    Mapea los encabezados del Excel a los campos requeridos del sistema
    usando una estrategia de dos fases (primero coincidencia exacta normalizada,
    luego coincidencia por inclusión) para evitar colisiones.
    """
    mapping = {}
    normalized_headers = [normalize_text(h) for h in raw_headers]
    used_indices = set()

    # Fase 1: Coincidencia EXACTA
    for field, patterns in COLUMN_PATTERNS.items():
        for p in patterns:
            for idx, norm_h in enumerate(normalized_headers):
                if idx not in used_indices and norm_h == p:
                    mapping[field] = idx
                    used_indices.add(idx)
                    break
            if field in mapping:
                break

    # Fase 2: Coincidencia POR CONTENCIÓN (para los campos que falten)
    for field, patterns in COLUMN_PATTERNS.items():
        if field in mapping:
            continue
        for p in patterns:
            for idx, norm_h in enumerate(normalized_headers):
                if idx not in used_indices and (p in norm_h or norm_h.startswith(p)):
                    mapping[field] = idx
                    used_indices.add(idx)
                    break
            if field in mapping:
                break

    return mapping

def parse_excel_file(file_path: str) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
    """
    Lee un archivo Excel (.xlsx o .xls), detecta automáticamente las columnas
    y retorna la lista de prendas normalizadas junto con estadísticas básicas.
    """
    ext = os.path.splitext(file_path)[1].lower()
    items = []

    # Intentamos leer primero con openpyxl para .xlsx
    if ext == '.xlsx':
        wb = openpyxl.load_workbook(file_path, data_only=True)
        sheet = wb.active

        # Buscar la fila de encabezados (por lo general está en la fila 1, pero buscamos en las primeras 10)
        header_row_idx = None
        raw_headers = []
        
        for row_idx, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            if row_idx > 10:
                break
            # Verificamos si al menos una celda parece encabezado
            text_cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if len(text_cells) >= 3:
                norm_candidates = [normalize_text(c) for c in text_cells]
                # Si encontramos palabras clave de inventario
                if any(k in norm_candidates for k in ['referencia', 'ref', 'talla', 'color', 'descripcion', 'tienda', 'bodega', 'teorico', 'diferencia', 'ean']):
                    header_row_idx = row_idx
                    raw_headers = [str(c).strip() if c is not None else "" for c in row]
                    break

        if header_row_idx is None:
            # Asumir la fila 1 por defecto si no hubo match explícito
            header_row_idx = 1
            for row in sheet.iter_rows(values_only=True, max_row=1):
                raw_headers = [str(c).strip() if c is not None else "" for c in row]

        col_map = map_headers(raw_headers)

        # Leer filas de datos a partir de header_row_idx + 1
        for row_idx, row in enumerate(sheet.iter_rows(min_row=header_row_idx + 1, values_only=True), start=header_row_idx + 1):
            if not any(row):
                continue
            
            def get_val(field: str, default="") -> Any:
                idx = col_map.get(field)
                if idx is not None and idx < len(row):
                    val = row[idx]
                    return val if val is not None else default
                return default

            def to_int(val: Any) -> int:
                try:
                    if val is None or val == "":
                        return 0
                    return int(float(val))
                except (ValueError, TypeError):
                    return 0

            ref = clean_reference(get_val('reference'))
            name = str(get_val('name')).strip()
            barcode = clean_barcode(get_val('barcode'))

            # Evitar filas vacías o de totales al final
            if not ref and not name and not barcode:
                continue
            if 'total' in ref.lower() or 'total' in name.lower():
                continue

            size = str(get_val('size')).strip()
            color = str(get_val('color')).strip()
            store = to_int(get_val('store_count'))
            warehouse = to_int(get_val('warehouse_count'))
            theoretical = to_int(get_val('theoretical_count'))
            
            # Si el Excel traía la diferencia, la usamos; si no, la calculamos: (Tienda + Bodega) - Teórico
            if 'difference' in col_map:
                diff = to_int(get_val('difference'))
            else:
                diff = (store + warehouse) - theoretical

            # Decodificación oficial Seven Seven (28 Dama, 45 Caballero, y tipo según código de 2 dígitos)
            gender, garment_code, garment_type = decode_seven_seven_reference(ref, name)

            # Categoría directa si existe en el Excel, si no el tipo oficial decodificado o inferido
            cat_val = str(get_val('category')).strip() if 'category' in col_map else ""
            if cat_val and cat_val != '-' and cat_val != '0':
                category = cat_val.title()
            elif garment_type:
                category = garment_type
            else:
                category = classify_category(name)

            obs_val = str(get_val('observation')).strip() if 'observation' in col_map else ""

            items.append({
                'reference': ref if ref else barcode,
                'name': name if name else 'Prenda Sin Nombre',
                'size': size if size else '-',
                'color': color if color else '-',
                'barcode': barcode if barcode else '-',
                'store_count': store,
                'warehouse_count': warehouse,
                'theoretical_count': theoretical,
                'difference': diff,
                'category': category,
                'gender': gender,
                'garment_code': garment_code,
                'garment_type': garment_type if garment_type else category,
                'observation': obs_val
            })

    else:
        # Fallback a pandas para .xls u otros formatos
        df = pd.read_excel(file_path)
        raw_headers = list(df.columns)
        col_map = map_headers(raw_headers)

        for _, row in df.iterrows():
            def get_df_val(field: str, default="") -> Any:
                idx = col_map.get(field)
                if idx is not None and idx < len(raw_headers):
                    val = row.iloc[idx]
                    return val if pd.notna(val) else default
                return default

            def to_int_df(val: Any) -> int:
                try:
                    if pd.isna(val) or val == "":
                        return 0
                    return int(float(val))
                except (ValueError, TypeError):
                    return 0

            ref = clean_reference(get_df_val('reference'))
            name = str(get_df_val('name')).strip()
            barcode = clean_barcode(get_df_val('barcode'))

            if not ref and not name and not barcode:
                continue
            if 'total' in ref.lower() or 'total' in name.lower():
                continue

            size = str(get_df_val('size')).strip()
            color = str(get_df_val('color')).strip()
            store = to_int_df(get_df_val('store_count'))
            warehouse = to_int_df(get_df_val('warehouse_count'))
            theoretical = to_int_df(get_df_val('theoretical_count'))

            if 'difference' in col_map:
                diff = to_int_df(get_df_val('difference'))
            else:
                diff = (store + warehouse) - theoretical

            # Decodificación oficial Seven Seven (28 Dama, 45 Caballero, y tipo según código de 2 dígitos)
            gender, garment_code, garment_type = decode_seven_seven_reference(ref, name)

            cat_val = str(get_df_val('category')).strip() if 'category' in col_map else ""
            if cat_val and cat_val != '-' and cat_val != '0':
                category = cat_val.title()
            elif garment_type:
                category = garment_type
            else:
                category = classify_category(name)

            obs_val = str(get_df_val('observation')).strip() if 'observation' in col_map else ""

            items.append({
                'reference': ref if ref else barcode,
                'name': name if name else 'Prenda Sin Nombre',
                'size': size if size else '-',
                'color': color if color else '-',
                'barcode': barcode if barcode else '-',
                'store_count': store,
                'warehouse_count': warehouse,
                'theoretical_count': theoretical,
                'difference': diff,
                'category': category,
                'gender': gender,
                'garment_code': garment_code,
                'garment_type': garment_type if garment_type else category,
                'observation': obs_val
            })

    meta = {
        'total_parsed': len(items),
        'columns_detected': list(col_map.keys()),
        'missing_columns': [f for f in COLUMN_PATTERNS.keys() if f not in col_map]
    }

    return items, meta
